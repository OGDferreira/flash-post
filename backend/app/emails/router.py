import csv
import io
import logging
import unicodedata
import uuid
from zipfile import BadZipFile

import httpx
from fastapi import APIRouter, Depends, HTTPException, Query, Request, Response, status
from fastapi.responses import RedirectResponse
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import case, func, select
from sqlalchemy.exc import SQLAlchemyError

from app.auth.dependencies import (
    AuthenticatedUser,
    DbSession,
    OwnerAccess,
    WorkspaceMemberAccess,
    require_csrf,
)
from app.core.crypto import decrypt_value, encrypt_value
from app.core.security import WorkspaceRole
from app.instagram.storage import SupabaseStorage, SupabaseStorageError
from app.models import EmailAccount
from app.schemas.emails import (
    EmailAccountAttachmentResponse,
    EmailAccountCounts,
    EmailAccountCreateRequest,
    EmailAccountItem,
    EmailAccountObservationRequest,
    EmailAccountStatusRequest,
    EmailAccountUpdateRequest,
    EmailAccountsImportResponse,
    EmailAccountsResponse,
)

router = APIRouter(prefix="/api/emails", tags=["Email account control"])
logger = logging.getLogger(__name__)
MAX_ATTACHMENT_BYTES = 10 * 1024 * 1024
MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 2000
SUPPORTED_IMAGE_TYPES = {
    "image/jpeg": ("jpg", lambda content: content.startswith(b"\xff\xd8\xff")),
    "image/png": ("png", lambda content: content.startswith(b"\x89PNG\r\n\x1a\n")),
    "image/webp": (
        "webp",
        lambda content: len(content) >= 12
        and content[0:4] == b"RIFF"
        and content[8:12] == b"WEBP",
    ),
}
STATUS_NEXT = {
    "available": "in_use",
    "in_use": "completed",
    "completed": "error",
    "error": "returned",
    "returned": "available",
}


def _storage_or_http_error() -> SupabaseStorage:
    try:
        return SupabaseStorage.from_settings()
    except RuntimeError:
        logger.error("Email error attachment storage is not configured.")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="O armazenamento privado de anexos não está configurado.",
        ) from None


def _response(account: EmailAccount) -> EmailAccountItem:
    try:
        password = decrypt_value(account.encrypted_password)
        two_factor_code = decrypt_value(account.encrypted_two_factor_code)
        two_factor_password = (
            decrypt_value(account.encrypted_two_factor_password)
            if account.encrypted_two_factor_password
            else ""
        )
    except ValueError as exc:
        logger.error(
            "Email account credentials could not be decrypted (account_id=%s).",
            account.id,
        )
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="As credenciais desta conta não puderam ser descriptografadas.",
        ) from exc
    return EmailAccountItem(
        id=account.id,
        supplier=account.supplier,
        email=account.email,
        password=password,
        responsible=account.responsible,
        status=account.status,
        observation=account.observation,
        two_factor_code=two_factor_code,
        two_factor_password=two_factor_password,
        attachment_url=(
            f"/api/emails/{account.id}/attachment"
            if account.error_attachment_path
            else None
        ),
        created_at=account.created_at,
        updated_at=account.updated_at,
    )


def _set_private_headers(response: Response) -> None:
    response.headers["Cache-Control"] = "no-store"
    response.headers["Pragma"] = "no-cache"


async def _account_for_workspace(
    account_id: uuid.UUID,
    workspace_id: uuid.UUID,
    db: DbSession,
    *,
    lock: bool = False,
) -> EmailAccount:
    query = select(EmailAccount).where(
        EmailAccount.id == account_id,
        EmailAccount.workspace_id == workspace_id,
    )
    if lock:
        query = query.with_for_update()
    account = await db.scalar(query)
    if account is None:
        raise HTTPException(status_code=404, detail="Conta de e-mail não encontrada.")
    return account


@router.get("", response_model=EmailAccountsResponse)
async def list_email_accounts(
    access: WorkspaceMemberAccess,
    db: DbSession,
    response: Response,
) -> EmailAccountsResponse:
    response.headers["Cache-Control"] = "no-store"
    accounts = (
        await db.scalars(
            select(EmailAccount)
            .where(EmailAccount.workspace_id == access.workspace.id)
            .order_by(EmailAccount.created_at.desc(), EmailAccount.id)
        )
    ).all()
    counts_row = (
        await db.execute(
            select(
                func.count(EmailAccount.id),
                func.sum(case((EmailAccount.status == "available", 1), else_=0)),
                func.sum(case((EmailAccount.status == "in_use", 1), else_=0)),
                func.sum(case((EmailAccount.status == "completed", 1), else_=0)),
                func.sum(case((EmailAccount.status == "error", 1), else_=0)),
                func.sum(case((EmailAccount.status == "returned", 1), else_=0)),
            ).where(EmailAccount.workspace_id == access.workspace.id)
        )
    ).one()
    return EmailAccountsResponse(
        can_manage=access.membership.role == WorkspaceRole.OWNER.value,
        accounts=[_response(account) for account in accounts],
        counts=EmailAccountCounts(
            total=counts_row[0] or 0,
            available=counts_row[1] or 0,
            in_use=counts_row[2] or 0,
            completed=counts_row[3] or 0,
            error=counts_row[4] or 0,
            returned=counts_row[5] or 0,
        ),
    )


def _normalize_import_header(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return " ".join(
        "".join(char for char in text if not unicodedata.combining(char))
        .strip()
        .casefold()
        .replace("-", "")
        .split()
    )


def _read_import_rows(filename: str, content: bytes) -> list[dict[str, str]]:
    if filename.casefold().endswith((".xlsx", ".xlsm")):
        try:
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=True
            )
        except (InvalidFileException, OSError, ValueError, KeyError, BadZipFile):
            raise HTTPException(400, "O arquivo Excel está inválido ou corrompido.") from None
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = next(values, None)
        if headers is None:
            workbook.close()
            return []
        rows = list(values)
        workbook.close()
    elif filename.casefold().endswith(".csv"):
        try:
            sample = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise HTTPException(400, "O arquivo CSV deve usar codificação UTF-8.") from exc
        try:
            dialect = csv.Sniffer().sniff(sample[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(sample), dialect)
        headers = next(reader, None)
        rows = list(reader)
    else:
        raise HTTPException(400, "Envie um arquivo CSV ou XLSX.")

    if not headers:
        raise HTTPException(400, "A planilha não contém a linha de cabeçalho.")
    indexes = {
        _normalize_import_header(header): index
        for index, header in enumerate(headers)
        if header is not None and str(header).strip()
    }
    required = {
        "fornecedor": "supplier",
        "email": "email",
        "senha": "password",
        "codigo 2fa": "two_factor_code",
        "senha do 2fa": "two_factor_password",
    }
    missing = [header for header in required if header not in indexes]
    if missing:
        raise HTTPException(
            400,
            "Colunas obrigatórias ausentes: " + ", ".join(missing) + ".",
        )

    parsed: list[dict[str, str]] = []
    for line_number, row in enumerate(rows, start=2):
        if not row or not any(str(value or "").strip() for value in row):
            continue
        if len(parsed) >= MAX_IMPORT_ROWS:
            raise HTTPException(400, f"O limite é de {MAX_IMPORT_ROWS} linhas por importação.")
        fields = {
            key: str(row[index] or "").strip() if index < len(row) else ""
            for header, key in required.items()
            for index in [indexes[header]]
        }
        try:
            validated = EmailAccountCreateRequest.model_validate(fields)
            if not validated.two_factor_password:
                raise ValueError("senha do 2FA vazia")
        except (ValueError, TypeError):
            raise HTTPException(
                400,
                f"Dados inválidos na linha {line_number}; revise fornecedor, e-mail e credenciais obrigatórias.",
            ) from None
        fields["email"] = str(validated.email)
        fields["supplier"] = validated.supplier
        parsed.append(fields)
    if not parsed:
        raise HTTPException(400, "A planilha não contém linhas válidas para importar.")
    emails = [row["email"] for row in parsed]
    if len(emails) != len(set(emails)):
        raise HTTPException(400, "A planilha contém e-mails repetidos.")
    return parsed


@router.post(
    "/import",
    response_model=EmailAccountsImportResponse,
    dependencies=[Depends(require_csrf)],
)
async def import_email_accounts(
    request: Request,
    access: OwnerAccess,
    db: DbSession,
    response: Response,
    filename: str = Query(..., min_length=1, max_length=255),
) -> EmailAccountsImportResponse:
    _set_private_headers(response)
    chunks: list[bytes] = []
    total_bytes = 0
    async for chunk in request.stream():
        total_bytes += len(chunk)
        if total_bytes > MAX_IMPORT_BYTES:
            raise HTTPException(413, "O arquivo deve ter no máximo 5 MB.")
        chunks.append(chunk)
    content = b"".join(chunks)
    try:
        rows = _read_import_rows(filename, content)
        existing_emails = (
            await db.scalars(
                select(EmailAccount.email).where(
                    EmailAccount.workspace_id == access.workspace.id,
                    func.lower(EmailAccount.email).in_(
                        [row["email"] for row in rows]
                    ),
                )
            )
        ).all()
        if existing_emails:
            raise HTTPException(
                status_code=409,
                detail="Uma ou mais contas desta planilha já existem neste workspace.",
            )
        accounts = [
            EmailAccount(
                workspace_id=access.workspace.id,
                supplier=row["supplier"],
                email=row["email"],
                encrypted_password=encrypt_value(row["password"]),
                encrypted_two_factor_code=encrypt_value(row["two_factor_code"]),
                encrypted_two_factor_password=encrypt_value(row["two_factor_password"]),
                status="available",
            )
            for row in rows
        ]
    except RuntimeError as exc:
        logger.error("Email credentials encryption is not configured (%s).", type(exc).__name__)
        raise HTTPException(
            status_code=503,
            detail="O armazenamento seguro de credenciais não está configurado.",
        ) from None
    db.add_all(accounts)
    try:
        await db.commit()
    except SQLAlchemyError:
        await db.rollback()
        raise HTTPException(
            status_code=500,
            detail="Não foi possível importar as contas de e-mail.",
        ) from None
    return EmailAccountsImportResponse(imported_count=len(accounts))


@router.post(
    "",
    response_model=EmailAccountItem,
    status_code=status.HTTP_201_CREATED,
    dependencies=[Depends(require_csrf)],
)
async def create_email_account(
    payload: EmailAccountCreateRequest,
    access: OwnerAccess,
    db: DbSession,
    response: Response,
) -> EmailAccountItem:
    _set_private_headers(response)
    try:
        account = EmailAccount(
            workspace_id=access.workspace.id,
            supplier=payload.supplier,
            email=str(payload.email),
            encrypted_password=encrypt_value(payload.password),
            encrypted_two_factor_code=encrypt_value(payload.two_factor_code),
            encrypted_two_factor_password=(
                encrypt_value(payload.two_factor_password)
                if payload.two_factor_password
                else None
            ),
            status="available",
        )
    except RuntimeError as exc:
        logger.error("Email credentials encryption is not configured (%s).", type(exc).__name__)
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="O armazenamento seguro de credenciais não está configurado.",
        ) from None
    db.add(account)
    try:
        await db.commit()
        await db.refresh(account)
    except SQLAlchemyError:
        await db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Não foi possível salvar a conta de e-mail.",
        ) from None
    return _response(account)


@router.put(
    "/{account_id}",
    response_model=EmailAccountItem,
    dependencies=[Depends(require_csrf)],
)
async def update_email_account(
    account_id: uuid.UUID,
    payload: EmailAccountUpdateRequest,
    access: OwnerAccess,
    db: DbSession,
    response: Response,
) -> EmailAccountItem:
    _set_private_headers(response)
    account = await _account_for_workspace(
        account_id, access.workspace.id, db, lock=True
    )
    try:
        account.supplier = payload.supplier
        account.email = str(payload.email)
        account.encrypted_password = encrypt_value(payload.password)
        account.encrypted_two_factor_code = encrypt_value(payload.two_factor_code)
        account.encrypted_two_factor_password = (
            encrypt_value(payload.two_factor_password)
            if payload.two_factor_password
            else None
        )
        await db.commit()
        await db.refresh(account)
    except RuntimeError:
        await db.rollback()
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="O armazenamento seguro de credenciais não está configurado.",
        ) from None
    except SQLAlchemyError:
        await db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Não foi possível atualizar a conta de e-mail.",
        ) from None
    return _response(account)


@router.patch(
    "/{account_id}/status",
    response_model=EmailAccountItem,
    dependencies=[Depends(require_csrf)],
)
async def advance_email_account_status(
    account_id: uuid.UUID,
    payload: EmailAccountStatusRequest,
    access: WorkspaceMemberAccess,
    user: AuthenticatedUser,
    db: DbSession,
    response: Response,
) -> EmailAccountItem:
    _set_private_headers(response)
    account = await _account_for_workspace(
        account_id, access.workspace.id, db, lock=True
    )
    next_status = STATUS_NEXT[account.status]
    if payload.status != next_status:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="O status deve avançar para a próxima etapa do ciclo.",
        )
    account.status = next_status
    account.responsible = user.nickname or user.full_name
    await db.commit()
    await db.refresh(account)
    return _response(account)


@router.patch(
    "/{account_id}/observation",
    response_model=EmailAccountItem,
    dependencies=[Depends(require_csrf)],
)
async def update_email_account_observation(
    account_id: uuid.UUID,
    payload: EmailAccountObservationRequest,
    access: WorkspaceMemberAccess,
    db: DbSession,
    response: Response,
) -> EmailAccountItem:
    _set_private_headers(response)
    account = await _account_for_workspace(
        account_id, access.workspace.id, db, lock=True
    )
    account.observation = payload.observation
    await db.commit()
    await db.refresh(account)
    return _response(account)


@router.post(
    "/{account_id}/attachment",
    response_model=EmailAccountAttachmentResponse,
    dependencies=[Depends(require_csrf)],
)
async def upload_email_error_attachment(
    account_id: uuid.UUID,
    request: Request,
    access: WorkspaceMemberAccess,
    db: DbSession,
    response: Response,
    filename: str = Query(min_length=1, max_length=255),
) -> EmailAccountAttachmentResponse:
    _set_private_headers(response)
    account = await _account_for_workspace(
        account_id, access.workspace.id, db, lock=True
    )
    content_type = (
        request.headers.get("content-type", "").split(";", 1)[0].strip().lower()
    )
    image = SUPPORTED_IMAGE_TYPES.get(content_type)
    if image is None:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail="Anexe uma imagem JPEG, PNG ou WebP.",
        )
    content_length = request.headers.get("content-length")
    if content_length:
        try:
            if int(content_length) > MAX_ATTACHMENT_BYTES:
                raise HTTPException(
                    status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                    detail="A imagem deve ter no máximo 10 MB.",
                )
        except ValueError:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Não foi possível verificar o tamanho do anexo.",
            ) from None
    chunks: list[bytes] = []
    size_bytes = 0
    async for chunk in request.stream():
        size_bytes += len(chunk)
        if size_bytes > MAX_ATTACHMENT_BYTES:
            raise HTTPException(
                status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                detail="A imagem deve ter no máximo 10 MB.",
            )
        chunks.append(chunk)
    content = b"".join(chunks)
    if not content or not image[1](content):
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail="O conteúdo enviado não corresponde a uma imagem válida.",
        )
    if not filename.strip():
        raise HTTPException(status_code=422, detail="Nome de arquivo inválido.")

    storage = _storage_or_http_error()
    path = (
        f"{access.workspace.id}/email-errors/{account.id}/"
        f"{uuid.uuid4()}.{image[0]}"
    )
    try:
        await storage.upload(path, content, content_type)
    except (SupabaseStorageError, httpx.HTTPError) as exc:
        logger.error("Email error attachment upload failed (%s).", type(exc).__name__)
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail="Não foi possível salvar a imagem de erro.",
        ) from None

    previous_path = account.error_attachment_path
    account.error_attachment_path = path
    try:
        await db.commit()
        await db.refresh(account)
    except SQLAlchemyError:
        await db.rollback()
        try:
            await storage.delete(path)
        except (SupabaseStorageError, httpx.HTTPError) as cleanup_error:
            logger.error(
                "Orphaned email attachment cleanup failed (%s).",
                type(cleanup_error).__name__,
            )
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Não foi possível vincular a imagem à conta.",
        ) from None
    if previous_path:
        try:
            await storage.delete(previous_path)
        except (SupabaseStorageError, httpx.HTTPError) as exc:
            logger.error(
                "Previous email attachment cleanup failed (%s).",
                type(exc).__name__,
            )
    return EmailAccountAttachmentResponse(
        attachment_url=f"/api/emails/{account.id}/attachment"
    )


@router.get("/{account_id}/attachment")
async def get_email_error_attachment(
    account_id: uuid.UUID,
    access: WorkspaceMemberAccess,
    db: DbSession,
) -> RedirectResponse:
    account = await _account_for_workspace(account_id, access.workspace.id, db)
    if not account.error_attachment_path:
        raise HTTPException(status_code=404, detail="Esta conta não possui anexo.")
    storage = _storage_or_http_error()
    try:
        signed_url = await storage.create_signed_url(account.error_attachment_path)
    except (SupabaseStorageError, httpx.HTTPError) as exc:
        logger.error("Email error attachment signing failed (%s).", type(exc).__name__)
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail="Não foi possível abrir o anexo da conta.",
        ) from None
    return RedirectResponse(signed_url, headers={"Cache-Control": "no-store"})


@router.delete(
    "/{account_id}",
    status_code=status.HTTP_204_NO_CONTENT,
    dependencies=[Depends(require_csrf)],
)
async def delete_email_account(
    account_id: uuid.UUID,
    access: OwnerAccess,
    db: DbSession,
) -> None:
    account = await _account_for_workspace(
        account_id, access.workspace.id, db, lock=True
    )
    if account.error_attachment_path:
        storage = _storage_or_http_error()
        try:
            await storage.delete(account.error_attachment_path)
        except (SupabaseStorageError, httpx.HTTPError) as exc:
            logger.error("Email attachment deletion failed (%s).", type(exc).__name__)
            raise HTTPException(
                status_code=status.HTTP_502_BAD_GATEWAY,
                detail="Não foi possível remover o anexo; a conta foi preservada.",
            ) from None
    await db.delete(account)
    await db.commit()
